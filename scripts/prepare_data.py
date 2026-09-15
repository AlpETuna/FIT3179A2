"""Reproducible, deliberately small extracts of the published source tables.

Run acquire_sources.py first. No synthetic households, coordinates or joint data.
"""
from pathlib import Path
import csv
import hashlib
import io
import json
import math
import zipfile
from datetime import datetime
import openpyxl

CACHE = Path('.cache/sources')
OUT = Path('public/data')
OUT.mkdir(parents=True, exist_ok=True)
STATES = {'1':'NSW','2':'Vic','3':'Qld','4':'SA','5':'WA','6':'Tas','7':'NT','8':'ACT'}

def write(name, value):
    (OUT / (name + '.json')).write_text(json.dumps(value, ensure_ascii=False, separators=(',', ':'), allow_nan=False), encoding='utf-8')

def table(geo, number):
    with zipfile.ZipFile(CACHE / f'census-{geo}.zip') as z:
        name = next(n for n in z.namelist() if f'_{number}_' in n and n.endswith('.csv'))
        return list(csv.DictReader(io.StringIO(z.read(name).decode('utf-8-sig'))))

def num(row, field):
    value = row[field]
    return int(value) if value not in ('..', '', 'np', 'N/A') else None

with zipfile.ZipFile(CACHE / 'census-AUS.zip') as z:
    path = next(n for n in z.namelist() if 'geog_desc' in n)
    w = openpyxl.load_workbook(io.BytesIO(z.read(path)), read_only=True, data_only=True)
    names = {str(r[1]):r[3] for r in w.worksheets[0].iter_rows(min_row=2, values_only=True) if r[1] is not None}

# G37's disjoint tenure totals; never include the rental-landlord subtotals.
tenures = [('Owned outright','O_OR_Total'),('Mortgage','O_MTG_Total'),('Renting','R_Tot_Total'),('Other tenure','Oth_ten_type_Total'),('Not stated','Ten_type_NS_Total')]
g37 = table('AUS','G37')[0]
total = num(g37,'Total_Total')
tenure = [dict(tenure=n,count=num(g37,f),share=num(g37,f)/total) for n,f in tenures]
write('tenure',tenure)
dwelling = [('Separate house','Sep_house'),('Townhouse','SemiD_ro_or_tce_h_th'),('Apartment','Flat_apart'),('Other dwelling','Oth_dwell'),('Not stated','not_stated')]
prefixes = ['O_OR_DS_','O_MTG_DS_','R_Tot_DS_','Oth_ten_type_DS_','Ten_type_NS_DS_']
mosaic=[]
for (label,suffix) in dwelling:
    for (tenure_name,_), prefix in zip(tenures,prefixes):
        field = prefix+suffix
        if label=='Townhouse' and prefix=='Oth_ten_type_DS_':field='Oth_ten_ty_DS_SD_ro_tce_h_th'
        if label=='Townhouse' and prefix=='Ten_type_NS_DS_':field='Ten_ty_NS_DS_SD_ro_tce_h_t'
        if label=='Apartment' and prefix=='Ten_type_NS_DS_':field='Ten_ty_NS_DS_Flat_apart'
        mosaic.append(dict(dwelling=label,tenure=tenure_name,count=num(g37,field)))
write('mosaic',mosaic)

g36=table('AUS','G36')[0]
tree=[dict(id='Homes',parent=None,name='Occupied private dwellings')]
leaves={
    'Separate house':[('Separate house','OPDs_Separate_house_Dwellings')],
    'Townhouse':[('One storey','OPDs_SD_r_t_h_th_1_sty_Dwgs'),('Two or more storeys','OPDs_SD_r_t_h_th_2_m_sty_Dwgs')],
    'Apartment':[('1–2 storey block','OPDs_F_ap_I_1or2_sty_blk_Ds'),('3 storey block','OPDs_F_ap_I_3_sty_blk_Dwgs'),('4–8 storey block','OPDs_F_ap_I_4to8_sty_blk_Ds'),('9+ storey block','OPDs_F_ap_I_9_m_sty_blk_Ds'),('Attached to a house','OPDs_Flt_apt_Att_house_Ds')],
    'Other dwelling':[('Caravan','OPDs_Oth_dwg_Cvn_Ds'),('Cabin or houseboat','OPDs_Oth_dwg_cab_hboat_Ds'),('Improvised home','OPDs_Ot_dwg_Im_hm_tnt_SO_Ds'),('Attached to shop/office','OPDs_Ot_dwg_Hs_f_att_sh_of_Ds')],
    'Not stated':[('Not stated','OPDs_Dwlling_structur_NS_Dwgs')]}
for category, fields in leaves.items():
    tree.append(dict(id=category,parent='Homes',name=category,category=category))
    for name, field in fields:
        tree.append(dict(id=category+'/'+name,parent=category,name=name,category=category,count=num(g36,field)))
write('dwellings',tree)

g42=table('AUS','G42')[0]
flows=[]
for d,p in [('Separate house','SH_'),('Townhouse','SDRTHT_Tot_'),('Apartment','FoA_Tot_'),('Other / not stated',None)]:
    for family, s in [('Family household','FHs_Tot'),('One person','Lone_P_H'),('Group household','Group_H')]:
        count=num(g42,p+s) if p else num(g42,'OD_Tot_'+s)+num(g42,'DS_NS_'+s)
        flows.append(dict(household=family,dwelling=d,count=count))
write('households',flows)

regions={}
excluded={}
for geo in ['SA3','SA4','STE']:
    rows=[]
    costs={r[f'{geo}_CODE_2021']:r for r in table(geo,'G02')}
    for r in table(geo,'G37'):
        code=r[f'{geo}_CODE_2021']; state=STATES.get(code[0])
        c=costs[code]; count=num(r,'Total_Total'); rent=num(c,'Median_rent_weekly')
        # Exclude migratory/offshore/no-usual-address codes and other territories.
        if not state or '9' == code[1:2] or not count or not rent:
            excluded[code]=names.get(code,code);continue
        rows.append(dict(code=code,name=names[code],state=state,households=count,
            renting=num(r,'R_Tot_Total'),rentalShare=num(r,'R_Tot_Total')/count,
            apartmentShare=num(r,'Total_DS_Flat_apart')/count,
            rent=rent,income=num(c,'Median_tot_hhd_inc_weekly'),
            mortgage=num(c,'Median_mortgage_repay_monthly'),
            mortgageWeekly=round(num(c,'Median_mortgage_repay_monthly')*12/52,2)))
    write(geo.lower(),rows);regions[geo]=rows
cases=[]
for name,kind in [('Melbourne City','Inner city'),('Casey - South','Outer suburb'),('Bendigo','Regional')]:
    row=next(r for r in regions['SA3'] if r['name']==name)
    cases.append({**row,'kind':kind})
write('cases',cases)

# Published monthly series. The two rent columns are duplicate sub-group/class
# series; choose the identified sub-group and assert they agree in this release.
w=openpyxl.load_workbook(CACHE/'cpi-03.xlsx',read_only=True,data_only=True)
rows=list(w['Data1'].values)
assert rows[9][1]=='A130393720C' and rows[9][55]=='A130390094A'
assert all(r[55]==r[56] for r in rows[10:])
common=[r for r in rows[10:] if isinstance(r[0],datetime) and isinstance(r[1],(int,float)) and isinstance(r[55],(int,float))]
cpi=[]
for r in common:
    for label,col in [('All prices',1),('Rents',55)]:
        cpi.append(dict(date=r[0].strftime('%Y-%m-%d'),series=label,sourceIndex=r[col],index=round(r[col]/common[0][col]*100,4)))
write('cpi',cpi)

# Cross-tabs exist in CRA.2. Use the published percentages, not rounded counts,
# for the rate. Denominators exclude units whose affordability cannot be calculated.
w=openpyxl.load_workbook(CACHE/'cra-quarterly.xlsx',read_only=True,data_only=True)
rows2=list(w['CRA.2'].values);rows3=list(w['CRA.3'].values)
date=max(r[0] for r in rows2 if isinstance(r[0],datetime))
current=[r for r in rows2 if r[0]==date and r[1]=='Per cent']
assert current, 'Do not calculate unpublished joint percentages'
def cra_row(description,category,detail):
    matches=[r for r in current if r[2].startswith(description) and r[3].startswith(category) and r[4]==detail]
    assert len(matches)==1,(description,category,detail,len(matches))
    return matches[0]
states=list(rows2[3][5:13])
before=cra_row('In rental stress excluding','Primary payment','Total')
after=cra_row('In rental stress including','Primary payment','Total')
severe50=next(r for r in rows3 if r[0]==date and r[1]=='Per cent' and r[2]=='>50 to 70% of income on rent')
severe70=next(r for r in rows3 if r[0]==date and r[1]=='Per cent' and r[2]=='>70% of income on rent')
tiles={'WA':(0,1),'NT':(1,0),'SA':(1,1),'Qld':(2,0),'NSW':(2,1),'Vic':(2,2),'ACT':(3,2),'Tas':(2,3)}
cra=[]
for i,state in enumerate(states):
    x,y=tiles[state]
    cra.append(dict(state=state,before=before[i+5],after=after[i+5],severe=severe50[i+3]+severe70[i+3],col=x,row=y))
write('cra',cra)
payments=[('Jobseeker(l)','JobSeeker'),('Parenting Payment (Single)','Single parent'),('Youth Allowance (Student & Apprentice)','Student / apprentice'),('Disability Support Pension(k)','Disability support'),('Age Pension','Age pension'),('Carer Payment','Carer')]
heat=[]
for full,short in payments:
    r=cra_row('In rental stress including','Primary payment',full)
    for i,state in enumerate(states):
        value=r[i+5]
        heat.append(dict(state=state,payment=short,fullPayment=full,rate=value if isinstance(value,(float,int)) else None))
write('heatmap',heat)

national=table('AUS','G02')[0]
write('summary',dict(households=total,rentalShare=num(g37,'R_Tot_Total')/total,
    nationalRent=num(national,'Median_rent_weekly'),nationalIncome=num(national,'Median_tot_hhd_inc_weekly'),
    craBefore=before[13],craAfter=after[13],craSevere=severe50[11]+severe70[11],
    craDate=date.strftime('%Y-%m-%d'),cpiStart=common[0][0].strftime('%Y-%m-%d'),cpiEnd=common[-1][0].strftime('%Y-%m-%d'),
    excluded=excluded,regionCount=len(regions['SA3'])))
write('provenance',dict(retrieved='2026-09-15',census='2021-08-10',cra=date.strftime('%Y-%m-%d'),
    files=[dict(file=p.name,bytes=p.stat().st_size,sha256=hashlib.sha256(p.read_bytes()).hexdigest()) for p in CACHE.iterdir() if p.suffix in ['.zip','.xlsx']],
    notes=['Census cell perturbation means components may not add to independently published totals.',
    'Composition layouts normalise disjoint cells; exact counts and denominator are supplied separately.',
    'CRA percentages are published CRA.2 and CRA.3 values, not recomputed from rounded counts.',
    'CPI index is rescaled to 100 at first common monthly observation, not spliced to the quarterly series.',
    'No Census regional median is treated as an individual household observation.']))
print(f'Prepared {len(list(OUT.glob("*.json")))} data files; {len(regions["SA3"])} SA3 regions, {len(regions["SA4"])} SA4 regions; {len(common)} CPI months; {len(heat)} verified heatmap cells.')
